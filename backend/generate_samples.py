import os
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document

SAMPLES_DIR = Path(__file__).resolve().parent.parent / "sample_resumes"
SAMPLES_DIR.mkdir(parents=True, exist_ok=True)

def generate_pdf_resume():
    file_path = SAMPLES_DIR / "sarah_chen_data_analyst.pdf"
    c = canvas.Canvas(str(file_path), pagesize=letter)
    width, height = letter
    y = height - 50

    c.setFont("Helvetica-Bold", 18)
    c.drawString(50, y, "Sarah Chen")
    y -= 20

    c.setFont("Helvetica", 10)
    c.drawString(50, y, "sarah.chen@university.edu | (555) 345-6789 | linkedin.com/in/sarahchen")
    y -= 30

    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "Education")
    y -= 18
    c.setFont("Helvetica", 11)
    c.drawString(50, y, "Bachelor of Science in Statistics & Data Science, University of California, 2025")
    y -= 25

    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "Technical Skills")
    y -= 18
    c.setFont("Helvetica", 11)
    c.drawString(50, y, "SQL (PostgreSQL), Python (Pandas, NumPy, Scikit-learn), Tableau, PowerBI, Excel, ETL, Git")
    y -= 25

    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "Experience")
    y -= 18
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "Business Intelligence Intern — TechCorp Solutions")
    y -= 16
    c.setFont("Helvetica", 10)
    c.drawString(60, y, "• Engineered automated SQL ETL pipelines, reducing query turnaround time by 40%.")
    y -= 14
    c.drawString(60, y, "• Designed interactive Tableau executive dashboards tracking $2.5M in annual customer spend.")
    y -= 14
    c.drawString(60, y, "• Cleaned and modeled 500k customer transaction rows to predict quarterly churn.")
    y -= 22

    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "Data Analyst Assistant — Campus Research Lab")
    y -= 16
    c.setFont("Helvetica", 10)
    c.drawString(60, y, "• Analyzed survey responses from 1,200 participants using Python and statistical hypothesis testing.")
    y -= 14
    c.drawString(60, y, "• Built predictive regression models improving forecast accuracy by 18%.")

    c.save()
    print(f"Generated {file_path}")

def generate_docx_resume():
    file_path = SAMPLES_DIR / "marcus_vance_consultant.docx"
    doc = Document()

    doc.add_heading("Marcus Vance", level=1)
    doc.add_paragraph("marcus.vance@mba.edu | (555) 789-0123 | linkedin.com/in/marcusvance")

    doc.add_heading("Education", level=2)
    doc.add_paragraph("Master of Business Administration (MBA), Kellogg School of Management, 2025")

    doc.add_heading("Core Competencies & Skills", level=2)
    doc.add_paragraph("Financial Modeling, Valuation, Strategy, PowerPoint, Excel, Stakeholder Management, Problem Solving")

    doc.add_heading("Professional Experience", level=2)
    p1 = doc.add_paragraph()
    p1.add_run("Strategy Consulting Intern — Apex Partners").bold = True
    doc.add_paragraph("• Formulated commercial due diligence strategy for a $45M healthcare acquisition.", style='List Bullet')
    doc.add_paragraph("• Spearheaded competitor benchmarking analysis across 14 European markets.", style='List Bullet')
    doc.add_paragraph("• Designed and presented 35-slide executive presentation decks directly to client CEO.", style='List Bullet')

    doc.add_heading("Leadership & Projects", level=2)
    doc.add_paragraph("President — Graduate Consulting Club", style='List Bullet')
    doc.add_paragraph("• Organized case competition for 180 MBA students with 5 sponsor consulting firms.", style='List Bullet')

    doc.save(str(file_path))
    print(f"Generated {file_path}")

def generate_scanned_pdf():
    # An empty PDF with no text elements (simulates pure raster/scanned image without text layer)
    file_path = SAMPLES_DIR / "scanned_image_only_resume.pdf"
    c = canvas.Canvas(str(file_path), pagesize=letter)
    # Draw blank rectangle / graphic only
    c.rect(50, 50, 500, 700, fill=0)
    c.save()
    print(f"Generated {file_path}")

if __name__ == "__main__":
    generate_pdf_resume()
    generate_docx_resume()
    generate_scanned_pdf()
