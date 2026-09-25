import pytest
from fastapi.testclient import TestClient
from main import app
from database import init_db, create_user, upgrade_user_to_pro, create_access_token
from services.resume_builder import CURATED_TEMPLATES

client = TestClient(app)

import tempfile

@pytest.fixture(autouse=True)
def setup_test_db(monkeypatch):
    temp_dir = tempfile.TemporaryDirectory()
    test_db_path = str(temp_dir.name + "/test_templates.db")
    monkeypatch.setattr("config.DATABASE_URL", None)
    monkeypatch.setattr("database.DATABASE_URL", None)
    monkeypatch.setattr("config.DATABASE_PATH", test_db_path)
    monkeypatch.setattr("database.DATABASE_PATH", test_db_path)
    init_db(test_db_path)
    user = create_user("Ivy Scholar", "ivy@college.edu", "dummyhash", db_path=test_db_path)
    upgrade_user_to_pro(user["id"], "CAMPUS2026", db_path=test_db_path)
    token = create_access_token(user["id"], user["email"])
    yield {"token": token, "db_path": test_db_path}
    temp_dir.cleanup()

SAMPLE_PAYLOAD = {
    "parsed_resume": {
        "candidate_name": "Alex Mercer",
        "contact_info": {
            "email": "alex@university.edu",
            "phone": "+1 (555) 019-2834",
            "linkedin": "linkedin.com/in/alexmercer",
            "github": "github.com/alexmercer"
        },
        "summary": "Computer Science student specializing in distributed systems and cloud platforms.",
        "skills": ["Python", "Go", "React", "PostgreSQL", "Docker", "AWS"],
        "education": [
            {
                "degree": "B.S. in Computer Science",
                "institution": "University of Technology",
                "year": "Expected May 2026"
            }
        ],
        "experience": [
            {
                "role": "Software Engineering Intern",
                "company": "Acme Cloud Corp",
                "duration": "June 2025 - August 2025",
                "bullets": [
                    "Helped team improve search database query times by adding indexes.",
                    "Built automated testing suite for backend API endpoints."
                ]
            }
        ],
        "projects": [
            {
                "title": "Distributed Task Scheduler",
                "tech_stack": "Go, Redis, Docker",
                "bullets": [
                    "Engineered worker pool handling 10,000 tasks/second with fault tolerance."
                ]
            }
        ]
    },
    "applied_rewrites": [
        {
            "original_bullet": "Helped team improve search database query times by adding indexes.",
            "rewrite_bullet": "Optimized database query performance by 42% by architecting composite B-tree indexes across PostgreSQL clusters.",
            "role": "Software Engineering Intern"
        }
    ]
}

def test_export_curated_templates_docx(setup_test_db):
    token = setup_test_db["token"]
    for template_id in ["ivy_league", "tech_minimalist", "modern_corporate"]:
        payload = {**SAMPLE_PAYLOAD, "format": "docx", "template_id": template_id}
        res = client.post("/api/resume/export", json=payload, headers={"Authorization": f"Bearer {token}"})
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert f"Alex_Mercer_{template_id}_Optimized.docx" in res.headers["content-disposition"]
        assert len(res.content) > 1000

def test_export_curated_templates_pdf(setup_test_db):
    token = setup_test_db["token"]
    for template_id in ["ivy_league", "tech_minimalist", "modern_corporate"]:
        payload = {**SAMPLE_PAYLOAD, "format": "pdf", "template_id": template_id}
        res = client.post("/api/resume/export", json=payload, headers={"Authorization": f"Bearer {token}"})
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/pdf"
        assert f"Alex_Mercer_{template_id}_Optimized.pdf" in res.headers["content-disposition"]
        assert len(res.content) > 1000

def test_export_original_template_in_place_docx(setup_test_db):
    token = setup_test_db["token"]
    import io
    from docx import Document

    # Create dummy original docx
    orig_doc = Document()
    orig_doc.add_heading("Alex Mercer", 0)
    orig_doc.add_paragraph("Helped team improve search database query times by adding indexes.")
    buf = io.BytesIO()
    orig_doc.save(buf)
    orig_bytes = buf.getvalue()

    # Scan the document first to persist raw bytes
    scan_res = client.post(
        "/api/scan",
        files={"file": ("alex_mercer.docx", orig_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"mode": "general"}
    )
    assert scan_res.status_code == 200
    scan_id = scan_res.json()["scan_id"]

    # Export with template_id == 'original' using scan_id
    payload = {
        **SAMPLE_PAYLOAD,
        "format": "docx",
        "template_id": "original",
        "scan_id": scan_id
    }
    res = client.post("/api/resume/export", json=payload, headers={"Authorization": f"Bearer {token}"})
    assert res.status_code == 200
    assert "OriginalFormat_Optimized.docx" in res.headers["content-disposition"]
    exported_doc = Document(io.BytesIO(res.content))
    full_text = "\n".join(p.text for p in exported_doc.paragraphs)
    assert "Optimized database query performance by 42%" in full_text

def test_export_custom_docx_template(setup_test_db):
    import io, base64
    from docx import Document
    token = setup_test_db["token"]

    # Create dummy custom docx
    custom_doc = Document()
    custom_doc.add_heading("Custom College Format", 0)
    p = custom_doc.add_paragraph("Helped team improve search database query times by adding indexes.")
    buf = io.BytesIO()
    custom_doc.save(buf)
    b64_content = base64.b64encode(buf.getvalue()).decode("utf-8")

    payload = {
        **SAMPLE_PAYLOAD,
        "format": "docx",
        "template_id": "custom",
        "custom_template_base64": b64_content
    }
    res = client.post("/api/resume/export", json=payload, headers={"Authorization": f"Bearer {token}"})
    assert res.status_code == 200
    assert "CustomTemplate_Optimized.docx" in res.headers["content-disposition"]
    
    # Read output and verify bullet was replaced in place
    exported_doc = Document(io.BytesIO(res.content))
    full_text = "\n".join(p.text for p in exported_doc.paragraphs)
    assert "Optimized database query performance by 42%" in full_text

