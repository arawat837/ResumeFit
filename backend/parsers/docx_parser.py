import io
import logging
from typing import Tuple, Dict, Any
from docx import Document
from fastapi import HTTPException
from config import MIN_TEXT_CHARS

logger = logging.getLogger(__name__)

class DOCXParser:
    """
    Parses DOCX resumes using python-docx.
    Extracts text from paragraphs and tables, and computes layout metadata.
    Raises HTTPException(422) if extracted text is fewer than MIN_TEXT_CHARS.
    """

    @staticmethod
    def parse(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        try:
            doc = Document(io.BytesIO(file_bytes))
        except Exception as e:
            logger.error(f"Error opening DOCX: {e}")
            raise HTTPException(
                status_code=422,
                detail=f"Could not parse DOCX file. It may be corrupted or in an unsupported format: {str(e)}"
            )

        extracted_text_blocks = []
        has_tables = len(doc.tables) > 0
        paragraph_count = len(doc.paragraphs)

        # 1. Paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                extracted_text_blocks.append(text)

        # 2. Tables (if any, extract row content)
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    extracted_text_blocks.append(" | ".join(row_texts))

        full_text = "\n\n".join(extracted_text_blocks).strip()
        char_count = len(full_text)
        word_count = len(full_text.split())

        if char_count < MIN_TEXT_CHARS:
            logger.warning(f"DOCX has insufficient text ({char_count} chars)")
            raise HTTPException(
                status_code=422,
                detail="Could not extract meaningful text from this DOCX file. Please upload a valid resume."
            )

        formatting_meta = {
            "file_type": "docx",
            "page_count": max(1, word_count // 450),  # Estimated pages
            "char_count": char_count,
            "word_count": word_count,
            "has_tables": has_tables,
            "image_count": 0,
            "has_multi_column": False,
            "paragraph_count": paragraph_count,
        }

        return full_text, formatting_meta
